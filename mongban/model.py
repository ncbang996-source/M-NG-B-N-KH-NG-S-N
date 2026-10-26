"""Thông số móng: giá trị mặc định, đọc Excel (sheet TONG HOP), kiểm tra, khối lượng bê tông."""
from __future__ import annotations

import math
from dataclasses import dataclass, field

# ô Excel -> thông số (giống mb:cfg-cells của MONG_BAN_KHONG_SUON.lsp)
CELLS_NUM = {
    "L1": "K7", "L2": "K8", "B": "K10", "L": "K11", "H1": "K14", "H2": "K16",
    "HB": "K20", "HV": "K21", "BB": "K23", "HC": "K26", "SLOPE": "K27", "MPIT": "K28",
}
CELLS_STR = {"GD": "C2", "COT": "D3", "TENMONG": "B4", "BTM": "K2", "BTL": "K3", "CODE": "K4"}
CELL_HT = "C3"
CELLS_BLN = ("K29", "L29")
# cốt thép (mm)
CELLS_REBAR = {
    "RB_BOT_D": "C31", "RB_BOT_S": "D31", "RB_TOP_D": "C32", "RB_TOP_S": "D32",
    "RB_CH_D": "C33", "RB_CH_S": "D33", "RC_D": "C35", "RC_N": "D35",
    "RT_D": "C36", "RT_S": "D36", "RP_D": "C37", "RP_S": "D37", "RL_D": "C38", "RL_S": "D38",
}
CELLS_COVER = {"COVER_B": "K31", "COVER": "K32"}   # mm trong Excel -> m
CELL_RLAYERS = "F38"                               # "02 lớp"

HT_NAMES = ["ĐỠ", "NÉO GÓC", "NÉO DỪNG"]
SCALES = [25, 50, 75, 100, 150, 200, 250]
SLOPE_DRAW = 0.5          # hệ số mở mái THỂ HIỆN trên bản vẽ (K27 chỉ dùng tính khối lượng)


def default_params() -> dict:
    return {
        "L1": 2.6, "L2": 2.6, "B": 5.0, "L": 6.0, "H1": 2.5, "H2": 1.0,
        "HB": 0.5, "HV": 0.3, "BB": 1.0,
        "TL": 0.1,      # dày lót (m)
        "ML": 0.1,      # lót tràn ra mỗi bên (m)
        "MPL": 0.1,     # vát tràn ra ngoài mặt cổ móng (m)
        "HC": 0.5,      # dày đệm cát (m), 0 = không có
        "SLOPE": 1.0,   # hệ số mở mái hố đào (tính khối lượng)
        "MPIT": 0.3,    # mở rộng đáy hố đào tính từ mép đế móng (m)
        "HT": 1,        # 0 = Đỡ, 1 = Néo góc, 2 = Néo dừng
        "BLN": "BLN-M36-250", "GD": "TKBVTC", "TENMONG": "",
        "SCALE": 100,
        "BTM": "B20", "BTL": "B7.5",
        "COT": "D411-25", "CODE": "",
        "RB_BOT_D": 12.0, "RB_BOT_S": 200.0, "RB_TOP_D": 12.0, "RB_TOP_S": 200.0,
        "RB_CH_D": 12.0, "RB_CH_S": 600.0, "RC_D": 20.0, "RC_N": 36.0,
        "RT_D": 8.0, "RT_S": 200.0, "RP_D": 8.0, "RP_S": 400.0,
        "RL_D": 8.0, "RL_S": 100.0, "RL_N": 2.0,
        "COVER": 0.05, "COVER_B": 0.05,
        # khung tên
        "MABV": "2604HPLB.F.TR.C.02",
    }


def type_name(p: dict) -> str:
    if p.get("TENMONG"):
        return p["TENMONG"]
    f = lambda v: f"{v:g}"
    return f"MB {f(p['L1'])}-{f(p['B'])}x{f(p['L'])}"


def ht_from_string(s: str) -> int:
    u = s.upper()
    if "GÓC" in u or "GOC" in u:
        return 1
    if "DỪNG" in u or "DUNG" in u:
        return 2
    return 0


def read_excel(path: str, sheet: str = "TONG HOP", p: dict | None = None) -> tuple[dict, list[str]]:
    """Đọc sheet TONG HOP. Trả về (thông số, danh sách cảnh báo)."""
    import openpyxl

    p = dict(p or default_params())
    warn: list[str] = []
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    if sheet not in wb.sheetnames:
        raise ValueError(f'Không tìm thấy sheet "{sheet}" trong file Excel.')
    ws = wb[sheet]
    val = lambda a: ws[a].value

    nread = 0
    for k, a in CELLS_NUM.items():
        v = val(a)
        if isinstance(v, (int, float)):
            p[k] = float(v)
            nread += 1
        elif k == "HC" and v in (None, ""):
            p["HC"] = 0.0
        else:
            warn.append(f"Ô {a} ({k}) không có số, giữ giá trị {p[k]}.")
    for k, a in CELLS_STR.items():
        v = val(a)
        if v not in (None, ""):
            p[k] = str(v).strip()
    v = val(CELL_HT)
    if isinstance(v, str):
        p["HT"] = ht_from_string(v)
    bln = "".join(str(val(a) or "").strip() for a in CELLS_BLN)
    if bln:
        p["BLN"] = bln
    for k, a in CELLS_REBAR.items():
        v = val(a)
        if isinstance(v, (int, float)) and v > 0:
            p[k] = float(v)
    for k, a in CELLS_COVER.items():
        v = val(a)
        if isinstance(v, (int, float)) and v > 0:
            p[k] = float(v) / 1000.0
    v = val(CELL_RLAYERS)
    if v not in (None, ""):
        try:
            n = int(v) if isinstance(v, (int, float)) else int("".join(c for c in str(v) if c.isdigit()) or 0)
            if n > 0:
                p["RL_N"] = float(n)
        except ValueError:
            pass
    wb.close()
    if nread == 0:
        raise ValueError("Không đọc được số liệu nào từ sheet. Kiểm tra lại địa chỉ ô.")
    p["TENMONG"] = p.get("TENMONG", "")
    return p, warn


def H(p):
    return p["H1"] + p["H2"]


def vol_mong(p) -> float:
    b, l, hv, hb, bb = p["B"], p["L"], p["HV"], p["HB"], p["BB"]
    px = p["L1"] + bb + 2 * p["MPL"]
    py = p["L2"] + bb + 2 * p["MPL"]
    return b * l * hv + (b * l + px * py + (b + py) * (l + px)) * (hb - hv) / 6.0 + 4.0 * bb * bb * (H(p) - hb)


def vol_lot(p) -> float:
    return (p["B"] + 2 * p["ML"]) * (p["L"] + 2 * p["ML"]) * p["TL"]


def vol_cat(p) -> float:
    hc = p["HC"]
    if hc <= 1e-6:
        return 0.0
    l, b, mp, s = p["L"], p["B"], p["MPIT"], p["SLOPE"]
    d = s * hc
    a1 = (l + 2 * mp) * (b + 2 * mp)
    a2 = (l + 2 * mp + 2 * d) * (b + 2 * mp + 2 * d)
    am = (l + 2 * mp + d) * (b + 2 * mp + d)
    return hc * (a1 + a2 + 4 * am) / 6.0


def area_vai(p) -> float:
    """Tấm vải trải phẳng hình chữ nhật: mỗi cạnh = đáy + 2 mái + 2 dải gấp + 1
    (dải gấp = mặt cát ngoài mép móng, +1 = phủ vào dưới móng 0.5 m mỗi bên); 4 góc tính đầy."""
    hc = p["HC"]
    if hc <= 1e-6:
        return 0.0
    l, b, mp, s = p["L"], p["B"], p["MPIT"], p["SLOPE"]
    d = s * hc
    ls = hc * math.sqrt(1 + s * s)
    fold = mp + d
    return (l + 2 * mp + 2 * ls + 2 * fold + 1) * (b + 2 * mp + 2 * ls + 2 * fold + 1)


def bolt_text(p, npier=4) -> str:
    s = p["BLN"]
    for i, c in enumerate(s):
        if c == "M" and i + 1 < len(s) and s[i + 1].isdigit():
            d = ""
            for ch in s[i + 1:]:
                if not ch.isdigit():
                    break
                d += ch
            return f"{4 * npier}M{d}"
    return s


def validate(p) -> str | None:
    bb = p["BB"]
    if min(p[k] for k in ("L1", "L2", "B", "L", "H1", "H2", "HB", "HV", "BB", "TL")) < 1e-6:
        return "Có kích thước <= 0. Kiểm tra lại các ô nhập."
    if p["HC"] < 0:
        return "Dày đệm cát không được âm (nhập 0 nếu không có)."
    if p["SLOPE"] < 1e-6:
        return "Hệ số mở mái phải > 0."
    if p["MPIT"] < p["ML"]:
        return "Mở rộng đáy hố đào phải >= lót ra mỗi bên."
    if p["L"] < p["L1"] + bb:
        return "L (chiều dài móng) phải >= L1 + b."
    if p["B"] < p["L2"] + bb:
        return "B (chiều rộng móng) phải >= L2 + b."
    if p["HV"] > p["HB"]:
        return "hv (cao mép) không được lớn hơn hb (cao giữa)."
    if p["HB"] >= H(p):
        return "H = H1 + H2 phải lớn hơn hb."
    if p["HB"] >= p["H1"]:
        return "H1 (chiều sâu chôn) phải lớn hơn hb."
    keys = ("RB_BOT_D", "RB_BOT_S", "RB_TOP_D", "RB_TOP_S", "RB_CH_D", "RB_CH_S", "RC_D",
            "RT_D", "RT_S", "RP_D", "RP_S", "RL_D", "RL_S")
    if min(p[k] for k in keys) < 1e-6:
        return "Đường kính / khoảng cách thép phải > 0."
    n = int(round(p["RC_N"]))
    if n < 4 or n % 4:
        return "Số thanh thép chủ cổ móng phải là bội số của 4 (>= 4)."
    if p["RL_N"] < 1:
        return "Số lớp lưới thép đầu trụ phải >= 1."
    if p["COVER"] < 0 or p["COVER_B"] < 0:
        return "Lớp BT bảo vệ không được âm."
    return None


def auto_scale(p) -> int:
    k1 = 172.1 / (H(p) + p["B"] + 2 * p["ML"] + p["HC"])
    k2 = 256.0 / (p["L"] + 2 * p["MPIT"] + 2 * SLOPE_DRAW * (p["H1"] + p["TL"] + p["HC"]))
    k = min(k1, k2)
    for d in SCALES:
        if 1000.0 / d <= k:
            return d
    return SCALES[-1]


def is_detail(p) -> bool:
    return "BVTC" in str(p.get("GD", "")).upper()
