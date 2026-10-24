"""Xuất bảng khối lượng ra Excel, mọi khối lượng tính bằng công thức tham chiếu sheet THONG SO."""
from __future__ import annotations

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.workbook.defined_name import DefinedName

from . import model
from .geom import KGM, Geo

THIN = Side(style="thin", color="808080")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
HEAD = PatternFill("solid", fgColor="DDEBF7")
INPUT = PatternFill("solid", fgColor="FFF2CC")
B = Font(bold=True)

# (tên, ký hiệu, khoá thông số, hệ số đổi, đơn vị)
# Lưu ý: tên định danh Excel không phân biệt hoa/thường -> cổ móng là TS_bc (tránh trùng TS_B)
PARAMS = [
    ("Tim chân cột theo X", "L1", "L1", 1, "m"),
    ("Tim chân cột theo Y", "L2", "L2", 1, "m"),
    ("Chiều dài móng (theo X)", "L", "L", 1, "m"),
    ("Chiều rộng móng (theo Y)", "B", "B", 1, "m"),
    ("Cạnh cổ móng", "b", "BB", 1, "m"),
    ("Chiều sâu chôn móng", "H1", "H1", 1, "m"),
    ("Chiều cao trên mặt đất", "H2", "H2", 1, "m"),
    ("Chiều cao bản móng ở giữa", "hb", "HB", 1, "m"),
    ("Chiều cao bản móng ở mép", "hv", "HV", 1, "m"),
    ("Dày bê tông lót", "tl", "TL", 1, "m"),
    ("Lót tràn ra mỗi bên", "ml", "ML", 1, "m"),
    ("Vát tràn ngoài cổ móng", "mpl", "MPL", 1, "m"),
    ("Dày đệm cát", "hc", "HC", 1, "m"),
    ("Hệ số mở mái hố đào", "m", "SLOPE", 1, ""),
    ("Mở rộng đáy hố đào (mỗi bên)", "mp", "MPIT", 1, "m"),
    ("Số thân trụ", "n", None, 4, "trụ"),
]
NAME = {"L1": "TS_L1", "L2": "TS_L2", "L": "TS_L", "B": "TS_B", "BB": "TS_bc", "H1": "TS_H1", "H2": "TS_H2",
        "HB": "TS_hb", "HV": "TS_hv", "TL": "TS_tl", "ML": "TS_ml", "MPL": "TS_mpl", "HC": "TS_hc",
        "SLOPE": "TS_m", "MPIT": "TS_mp", None: "TS_n"}


def _cell(ws, r, c, v, fmt=None, bold=False, fill=None, align=None):
    x = ws.cell(row=r, column=c, value=v)
    x.border = BOX
    if fmt:
        x.number_format = fmt
    if bold:
        x.font = B
    if fill:
        x.fill = fill
    x.alignment = align or Alignment(vertical="center", wrap_text=True)
    return x


def export(p: dict, g: Geo, path: str):
    if not getattr(g, "kl", None):
        g.calc_kl()
    wb = Workbook()
    ws = wb.active
    ws.title = "THONG SO"
    ws.column_dimensions["A"].width = 6
    ws.column_dimensions["B"].width = 36
    ws.column_dimensions["C"].width = 10
    ws.column_dimensions["D"].width = 12
    ws.column_dimensions["E"].width = 8
    ws["A1"] = f"THÔNG SỐ MÓNG {model.type_name(p)} - CỘT {model.HT_NAMES[int(p['HT'])]} {p['COT']} ({p['GD']})"
    ws["A1"].font = Font(bold=True, size=12)
    ws["A2"] = "Ô nền vàng là số liệu đầu vào - sửa ở đây, sheet KHOI LUONG tự tính lại."
    ws["A2"].font = Font(italic=True, color="808080")
    r = 4
    for i, h in enumerate(["STT", "Thông số", "Ký hiệu", "Giá trị", "Đơn vị"], 1):
        _cell(ws, r, i, h, bold=True, fill=HEAD, align=Alignment(horizontal="center", vertical="center"))
    for k, (nm, sym, key, f, unit) in enumerate(PARAMS, 1):
        r += 1
        v = p[key] * f if key else f
        _cell(ws, r, 1, k, align=Alignment(horizontal="center"))
        _cell(ws, r, 2, nm)
        _cell(ws, r, 3, sym, align=Alignment(horizontal="center"))
        _cell(ws, r, 4, v, "0.000", fill=INPUT)
        _cell(ws, r, 5, unit, align=Alignment(horizontal="center"))
        wb.defined_names[NAME[key]] = DefinedName(NAME[key], attr_text=f"'THONG SO'!$D${r}")
    # bảng tra trọng lượng thép
    r += 2
    ws.cell(row=r, column=2, value="TRỌNG LƯỢNG THÉP (TCVN 1651-2018)").font = B
    r += 1
    _cell(ws, r, 2, "Đường kính (mm)", bold=True, fill=HEAD)
    _cell(ws, r, 3, "kg/m", bold=True, fill=HEAD)
    r0 = r + 1
    for d, w in sorted(KGM.items()):
        r += 1
        _cell(ws, r, 2, d)
        _cell(ws, r, 3, w, "0.000", fill=INPUT)
    wb.defined_names["TS_KGM"] = DefinedName("TS_KGM", attr_text=f"'THONG SO'!$B${r0}:$C${r}")

    # ------------------------------------------------------------------ KHOI LUONG
    wk = wb.create_sheet("KHOI LUONG")
    for col, w in zip("ABCDEFGH", (6, 44, 10, 70, 12, 9, 12, 12)):
        wk.column_dimensions[col].width = w
    wk["A1"] = f"BẢNG KHỐI LƯỢNG CHO 01 VỊ TRÍ MÓNG {model.type_name(p)} - CỘT {p['COT']}"
    wk["A1"].font = Font(bold=True, size=12)
    r = 3
    for i, h in enumerate(["STT", "Hạng mục", "Ký hiệu", "Công thức / diễn giải", "Khối lượng", "Đơn vị"], 1):
        _cell(wk, r, i, h, bold=True, fill=HEAD, align=Alignment(horizontal="center", vertical="center"))
    rows = {}

    def line(stt, name, sym, formula, desc, unit, bold=False):
        nonlocal r
        r += 1
        _cell(wk, r, 1, stt, align=Alignment(horizontal="center"))
        _cell(wk, r, 2, name, bold=bold)
        _cell(wk, r, 3, sym, align=Alignment(horizontal="center"))
        _cell(wk, r, 4, desc)
        _cell(wk, r, 5, formula, "0.00", bold=bold)
        _cell(wk, r, 6, unit, align=Alignment(horizontal="center"))
        if sym:
            rows[sym] = f"$E${r}"
        return f"$E${r}"

    def sec(t):
        nonlocal r
        r += 1
        c = _cell(wk, r, 1, t, bold=True, fill=HEAD)
        wk.merge_cells(start_row=r, start_column=1, end_row=r, end_column=6)
        c.alignment = Alignment(vertical="center")

    sec("A. KÍCH THƯỚC TRUNG GIAN")
    px = line("", "Mặt trên bản móng theo X", "px", "=TS_L1+TS_bc+2*TS_mpl", "L1 + b + 2·mpl", "m")
    py = line("", "Mặt trên bản móng theo Y", "py", "=TS_L2+TS_bc+2*TS_mpl", "L2 + b + 2·mpl", "m")
    hd = line("", "Chiều sâu hố đào", "Hd", "=TS_H1+TS_tl+TS_hc", "H1 + dày lót + dày đệm cát", "m")
    a1 = line("", "Đáy hố đào theo X", "a1", "=TS_L+2*TS_mp", "L + 2·mp", "m")
    b1 = line("", "Đáy hố đào theo Y", "b1", "=TS_B+2*TS_mp", "B + 2·mp", "m")
    a2 = line("", "Miệng hố đào theo X", "a2", f"={a1}+2*TS_m*{hd}", "a1 + 2·m·Hd", "m")
    b2 = line("", "Miệng hố đào theo Y", "b2", f"={b1}+2*TS_m*{hd}", "b1 + 2·m·Hd", "m")
    dc = line("", "Mở rộng mái trong lớp đệm cát", "dc", "=TS_m*TS_hc", "m · hc", "m")

    sec("B. BÊ TÔNG, ĐỆM CÁT, VẢI ĐỊA")
    vb = line(1, f"Bê tông móng {p['BTM']}", "Vbt",
              f"=TS_B*TS_L*TS_hv+(TS_B*TS_L+{px}*{py}+(TS_B+{py})*(TS_L+{px}))*(TS_hb-TS_hv)/6+TS_n*TS_bc^2*(TS_H1+TS_H2-TS_hb)",
              "B·L·hv + (B·L + px·py + (B+py)(L+px))·(hb−hv)/6 + n·b²·(H1+H2−hb)", "m3", True)
    vl = line(2, f"Bê tông lót {p['BTL']}", "Vlot", "=(TS_B+2*TS_ml)*(TS_L+2*TS_ml)*TS_tl",
              "(B + 2·ml)·(L + 2·ml)·tl", "m3", True)
    vc = line(3, "Đệm cát", "Vcat",
              f"=IF(TS_hc>0,TS_hc/6*({a1}*{b1}+({a1}+2*{dc})*({b1}+2*{dc})+4*({a1}+{dc})*({b1}+{dc})),0)",
              "Hình chóp cụt: hc/6·(A_đáy + A_trên + 4·A_giữa)", "m3", True)
    line(4, "Vải địa kỹ thuật", "Svai",
         f"=IF(TS_hc>0,{a1}*{b1}+2*TS_hc*SQRT(1+TS_m^2)*({a1}/2+{a1}/2+{dc}+{b1}/2+{b1}/2+{dc})"
         f"+4*(({a1}/2+{dc})*({b1}/2+{dc})-(TS_L/2-0.5)*(TS_B/2-0.5)),0)",
         "Đáy + mái hố trong lớp cát + phần gấp vào 500 mm trên mặt cát", "m2", True)

    sec("C. ĐẤT, VÁN KHUÔN")
    vd = line(5, "Đào đất hố móng", "Vdao",
              f"={hd}/6*({a1}*{b1}+{a2}*{b2}+({a1}+{a2})*({b1}+{b2}))",
              "Hình chóp cụt: Hd/6·(a1·b1 + a2·b2 + (a1+a2)(b1+b2))", "m3", True)
    vu = line("", "Thể tích móng nằm dưới mặt đất", "Vmd", f"={vb}-TS_n*TS_bc^2*TS_H2", "Vbt − n·b²·H2", "m3")
    vdp = line(6, "Đắp đất hoàn trả hố móng", "Vdap", f"={vd}-{vu}-{vl}-{vc}", "Vdao − Vmd − Vlot − Vcat", "m3", True)
    line(7, "Đất thừa (đắp ụ quanh cổ móng)", "Vthua", f"={vd}-{vdp}", "Vdao − Vdap", "m3", True)
    vk1 = line("", "Ván khuôn mép bản móng", "VK1", "=2*(TS_L+TS_B)*TS_hv", "2·(L + B)·hv", "m2")
    vk2 = line("", "Ván khuôn cổ móng", "VK2", "=TS_n*4*TS_bc*(TS_H1+TS_H2-TS_hb)", "n·4·b·(H1 + H2 − hb)", "m2")
    line(8, "Ván khuôn móng", "VK", f"={vk1}+{vk2}", "VK1 + VK2 (mặt vát không ghép ván khuôn)", "m2", True)

    sec("D. CỐT THÉP (số thanh và chiều dài do chương trình tính theo bố trí trên bản vẽ)")
    r += 1
    for i, h in enumerate(["Vị trí", "Hình dạng / ghi chú", "Ø (mm)", "Chiều dài 1 thanh (mm)", "Số lượng",
                           "kg/m", "Tổng dài (m)", "KL (kg)"], 1):
        _cell(wk, r, i, h, bold=True, fill=HEAD, align=Alignment(horizontal="center", vertical="center", wrap_text=True))
    desc = {1: "Lớp dưới theo Y", 2: "Lớp trên theo Y (uốn theo mặt vát)", 3: "Lớp dưới theo X",
            4: "Lớp trên theo X (uốn theo mặt vát)", 5: "Thép chống", 6: "Thép chủ cổ móng", 7: "Đai cổ móng",
            9: "Đai phụ thẳng", 10: "Đai bát giác", 11: "Lưới thép đầu trụ"}
    rs0 = r + 1
    for row in g.kl:
        r += 1
        q = str(row["qty"]).replace("x", "*")
        _cell(wk, r, 1, g.sq(row["pos"]), align=Alignment(horizontal="center"))
        _cell(wk, r, 2, desc.get(row["pos"], ""))
        _cell(wk, r, 3, int(row["dia"] + 0.5), fill=INPUT)
        _cell(wk, r, 4, row["len"], "0", fill=INPUT)
        _cell(wk, r, 5, f"={q}", "0")
        _cell(wk, r, 6, f"=IFERROR(VLOOKUP(C{r},TS_KGM,2,FALSE),C{r}^2/162)", "0.000")
        _cell(wk, r, 7, f"=D{r}/1000*E{r}", "0.00")
        _cell(wk, r, 8, f"=G{r}*F{r}", "0.00")
    rs1 = r
    r += 1
    _cell(wk, r, 2, "TỔNG CỘNG THÉP", bold=True)
    _cell(wk, r, 8, f"=SUM(H{rs0}:H{rs1})", "0.00", bold=True)
    r += 2
    _cell(wk, r, 2, "Phân loại theo đường kính", bold=True, fill=HEAD)
    _cell(wk, r, 3, "Ø (mm)", bold=True, fill=HEAD)
    _cell(wk, r, 7, "Tổng dài (m)", bold=True, fill=HEAD)
    _cell(wk, r, 8, "KL (kg)", bold=True, fill=HEAD)
    for d, _, _ in g.klc:
        r += 1
        _cell(wk, r, 2, f"Thép Ø{int(d + 0.5)}")
        _cell(wk, r, 3, int(d + 0.5))
        _cell(wk, r, 7, f"=SUMIF($C${rs0}:$C${rs1},C{r},$G${rs0}:$G${rs1})", "0.00")
        _cell(wk, r, 8, f"=SUMIF($C${rs0}:$C${rs1},C{r},$H${rs0}:$H${rs1})", "0.00")
    r += 2
    _cell(wk, r, 2, "Bu lông neo", bold=True)
    _cell(wk, r, 4, f"{model.bolt_text(p)}  ({p['BLN']})")
    wk.freeze_panes = "A4"
    wb.active = 1
    wb.save(path)
